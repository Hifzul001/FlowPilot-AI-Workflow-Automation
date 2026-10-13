"""Build executable workflow config from the recruiter workbook, preserving its wording."""
from __future__ import annotations

import json
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "data" / "assessment.xlsx"
DEST = ROOT / "data" / "workflows_config.xlsx"

# Step definitions are declarative. Tools are shared and implemented in app/tools.py.
DEFINITIONS = {
    "WF001": [{"tool": "inventory_restock", "args": {"threshold": "{{threshold}}"}, "save_as": "result"}],
    "WF002": [{"tool": "price_validation", "args": {"threshold_pct": "{{threshold_pct}}"}, "save_as": "result"}],
    "WF003": [{"tool": "vendor_file_process", "args": {"upload": "{{upload}}"}, "save_as": "result"}],
    "WF004": [{"tool": "product_content", "args": {"product": "{{product}}"}, "save_as": "result"}],
    "WF005": [{"tool": "order_lookup", "args": {"order_id": "{{order_id}}", "email": "{{email}}"}, "save_as": "result"}],
    "WF006": [{"tool": "duplicate_detection", "args": {}, "save_as": "result"}],
    "WF007": [{"tool": "campaign_brief", "args": {"goal": "{{goal}}", "products": "{{products}}", "audience": "{{audience}}", "promotion": "{{promotion}}", "start_date": "{{start_date}}", "end_date": "{{end_date}}"}, "save_as": "result"}],
    "WF008": [{"tool": "keyword_classification", "args": {"keywords": "{{keywords}}", "category": "{{category}}"}, "save_as": "result"}],
    "WF009": [{"tool": "task_assignment", "args": {"task": "{{task}}", "priority": "{{priority}}", "deadline": "{{deadline}}"}, "save_as": "result"}],
    "WF010": [{"tool": "performance_report", "args": {}, "save_as": "result"}],
}

ARGUMENTS = {
    "WF001": {"threshold": {"type": "number", "required": False}},
    "WF002": {"threshold_pct": {"type": "number", "required": False, "default": 10}},
    "WF003": {"upload": {"type": "file", "required": True}},
    "WF004": {"product": {"type": "object", "required": True}},
    "WF005": {"order_id": {"type": "string", "required": False}, "email": {"type": "string", "required": False}},
    "WF006": {},
    "WF007": {"goal": {"type": "string", "required": True}, "products": {"type": "list", "required": False}, "audience": {"type": "string", "required": False}, "promotion": {"type": "string", "required": False}, "start_date": {"type": "string", "required": True}, "end_date": {"type": "string", "required": True}},
    "WF008": {"keywords": {"type": "list", "required": True}, "category": {"type": "string", "required": False, "default": "General"}},
    "WF009": {"task": {"type": "string", "required": True}, "priority": {"type": "string", "required": False, "default": "normal"}, "deadline": {"type": "string", "required": False}},
    "WF010": {},
}

def main() -> None:
    source = openpyxl.load_workbook(SOURCE, data_only=True)
    out = openpyxl.Workbook()
    ws = out.active
    ws.title = "Workflows"
    headers = [str(c.value) for c in source["Workflows"][1]] + ["Arguments_JSON", "Executable_Steps_JSON"]
    ws.append(headers)
    for row in source["Workflows"].iter_rows(min_row=2, values_only=True):
        if not row[0]:
            continue
        wid = row[0]
        ws.append(list(row) + [json.dumps(ARGUMENTS[wid]), json.dumps(DEFINITIONS[wid])])
    q = out.create_sheet("Test_Questions")
    for row in source["Test_Questions"].iter_rows(values_only=True):
        q.append(list(row))
    for sheet in out.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for c in sheet[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="1F4E78")
        for col in sheet.columns:
            letter = col[0].column_letter
            sheet.column_dimensions[letter].width = min(64, max(14, max(len(str(x.value or "")) for x in col) + 2))
    out.save(DEST)
    print(f"Created {DEST}")

if __name__ == "__main__":
    main()
