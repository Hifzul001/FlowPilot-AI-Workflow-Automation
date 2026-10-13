from __future__ import annotations

import json
from pathlib import Path

import openpyxl

from app.models import Workflow


class WorkflowLoader:
    """Read workflow metadata and executable step definitions from an Excel file."""

    def __init__(self, path: str | Path):
        self.path = Path(path)

    def load(self) -> dict[str, Workflow]:
        workbook = openpyxl.load_workbook(self.path, data_only=True, read_only=True)
        if "Workflows" not in workbook.sheetnames:
            raise ValueError("Configuration workbook must contain a Workflows sheet")
        sheet = workbook["Workflows"]
        headers = [str(cell.value).strip() if cell.value is not None else "" for cell in sheet[1]]
        index = {name: pos for pos, name in enumerate(headers)}
        required = {"Workflow_ID", "Workflow_Name", "Trigger", "Inputs", "Steps", "Decision_Logic", "Tools_Required", "Expected_Output"}
        absent = required - index.keys()
        if absent:
            raise ValueError(f"Missing workflow columns: {', '.join(sorted(absent))}")

        def value(row, key, default=""):
            pos = index.get(key)
            return row[pos] if pos is not None and pos < len(row) and row[pos] is not None else default

        workflows: dict[str, Workflow] = {}
        for row in sheet.iter_rows(min_row=2, values_only=True):
            if not row or not row[index["Workflow_ID"]]:
                continue
            wid = str(value(row, "Workflow_ID")).strip()
            args = json.loads(value(row, "Arguments_JSON", "{}"))
            steps = json.loads(value(row, "Executable_Steps_JSON", "[]"))
            wf = Workflow(
                workflow_id=wid,
                name=str(value(row, "Workflow_Name")),
                trigger=str(value(row, "Trigger")),
                inputs=str(value(row, "Inputs")),
                steps=str(value(row, "Steps")),
                decision_logic=str(value(row, "Decision_Logic")),
                tools_required=str(value(row, "Tools_Required")),
                expected_output=str(value(row, "Expected_Output")),
                arguments=args,
                executable_steps=steps,
            )
            if not wf.executable_steps:
                raise ValueError(f"Workflow {wid} has no executable steps")
            workflows[wid] = wf
        workbook.close()
        if not workflows:
            raise ValueError("No workflows found in configuration")
        return workflows
